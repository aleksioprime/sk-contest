from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from tests.helpers import login_as

pytestmark = pytest.mark.usefixtures('redis_clean')

SLUG = 'skolstart-2027'


def _other(client: AsyncClient) -> AsyncClient:
    return AsyncClient(transport=client._transport, base_url=str(client.base_url), headers=client.headers)


async def _open_event(db, client, **form_changes) -> dict:
    """Организатор создаёт конкурс проектов с направлением и открывает регистрацию."""
    organizer = await login_as(db, client, email='org@skolkovo.example.org')
    organizer.person.is_employee = True
    await db.commit()
    event = (
        await client.post('/api/v1/events', json={'type': 'project_contest', 'title': 'SkolStart', 'slug': SLUG})
    ).json()
    url = f'/api/v1/events/{event["id"]}'
    await client.patch(url, json={'date_start': '2027-03-12'})
    await client.post(f'{url}/directions', json={'title': 'Инженерия'})
    if form_changes:
        form = (await client.get(f'{url}/form')).json()
        for key, value in form_changes.items():
            form[key] = {**form[key], **value} if isinstance(value, dict) else value
        assert (await client.put(f'{url}/form', json=form)).status_code == 200
    await client.post(f'{url}/status', json={'status': 'registration'})
    return (await client.get(url)).json()


def _payload(event, **extra) -> dict:
    return {
        'title': 'Умная теплица',
        'direction_id': event['directions'][0]['id'],
        'participation_form_id': event['participation_forms'][0]['id'],
        'answers': {'annotation': 'Автоматический полив по датчикам влажности'},
        'members': [
            {'role': 'participant', 'full_name': 'Иванова Анна', 'grade': '9'},
            {'role': 'supervisor', 'full_name': 'Петров Иван', 'email': 'petrov@example.org'},
        ],
        'consents': ['personal_data'],
        **extra,
    }


async def _submitted(applicant, event, **extra) -> dict:
    app = (await applicant.post(f'/api/v1/events/{SLUG}/applications')).json()
    await applicant.put(f'/api/v1/applications/{app["id"]}', json=_payload(event, **extra))
    response = await applicant.post(f'/api/v1/applications/{app["id"]}/submit')
    assert response.status_code == 200, response.text
    return response.json()


async def test_public_form_has_defaults(db, client):
    await _open_event(db, client)
    form = (await client.get(f'/api/v1/public/events/{SLUG}/form')).json()
    assert form['registration_open'] is True
    assert [f['key'] for f in form['form']['fields']] == ['annotation', 'work']
    assert form['form']['team']['require_supervisor'] is True
    assert form['form']['consents'][0]['key'] == 'personal_data'


async def test_full_application_cycle(db, client, queue):
    event = await _open_event(db, client)
    async with _other(client) as applicant:
        await login_as(db, applicant, email='anna@example.org', last_name='Иванова')
        draft = (await applicant.post(f'/api/v1/events/{SLUG}/applications')).json()
        assert draft['number'] == 1 and draft['status'] == 'draft'
        assert draft['members'][0]['full_name'] == 'Иванова Тест' and draft['can_submit'] is True

        incomplete = await applicant.post(f'/api/v1/applications/{draft["id"]}/submit')
        assert incomplete.status_code == 422
        errors = incomplete.json()['detail']['errors']
        assert {'title', 'direction_id', 'annotation', 'consent:personal_data', 'members'} <= set(errors)

        saved = await applicant.put(f'/api/v1/applications/{draft["id"]}', json=_payload(event))
        assert saved.status_code == 200 and saved.json()['consents'][0]['version'] == '1'
        submitted = (await applicant.post(f'/api/v1/applications/{draft["id"]}/submit')).json()
        assert submitted['status'] == 'submitted' and submitted['can_edit'] is False
        assert queue.jobs[-1][1][0] == 'anna@example.org'

        locked = await applicant.put(f'/api/v1/applications/{draft["id"]}', json=_payload(event, title='Новое'))
        assert locked.status_code == 422

        # организатор просит правки: без комментария нельзя
        no_comment = await client.post(f'/api/v1/applications/{draft["id"]}/decision', json={'status': 'needs_changes'})
        assert no_comment.status_code == 422
        await client.post(
            f'/api/v1/applications/{draft["id"]}/decision',
            json={'status': 'needs_changes', 'comment': 'Добавьте класс руководителя'},
        )
        assert 'Добавьте класс руководителя' in queue.jobs[-1][1][2]

        mine = (await applicant.get('/api/v1/my/applications')).json()
        assert mine[0]['status'] == 'needs_changes' and mine[0]['can_edit'] is True
        await applicant.put(f'/api/v1/applications/{draft["id"]}', json=_payload(event, title='Умная теплица 2.0'))
        assert (await applicant.post(f'/api/v1/applications/{draft["id"]}/submit')).json()['status'] == 'submitted'

    accepted = await client.post(f'/api/v1/applications/{draft["id"]}/decision', json={'status': 'accepted'})
    assert accepted.json()['status'] == 'accepted'
    listing = (await client.get(f'/api/v1/events/{event["id"]}/applications')).json()
    assert listing['counts'] == {'accepted': 1}
    assert listing['items'][0]['participants'] == ['Иванова Анна']


async def test_strangers_cannot_see_or_moderate(db, client):
    event = await _open_event(db, client)
    async with _other(client) as applicant, _other(client) as stranger:
        await login_as(db, applicant, email='anna@example.org')
        app = await _submitted(applicant, event)
        await login_as(db, stranger, email='stranger@example.org')
        assert (await stranger.get(f'/api/v1/applications/{app["id"]}')).status_code == 404
        assert (await stranger.get(f'/api/v1/events/{event["id"]}/applications')).status_code == 404
        denied = await stranger.post(f'/api/v1/applications/{app["id"]}/decision', json={'status': 'accepted'})
        assert denied.status_code == 404


async def test_files_are_private(db, client):
    event = await _open_event(db, client)
    async with _other(client) as applicant, _other(client) as stranger:
        await login_as(db, applicant, email='anna@example.org')
        upload = await applicant.post(
            '/api/v1/files', files={'file': ('работа.pdf', b'%PDF-1.4 test', 'application/pdf')}
        )
        assert upload.status_code == 201
        file_id = upload.json()['id']
        bad = await applicant.post('/api/v1/files', files={'file': ('x.exe', b'MZ', 'application/x-msdownload')})
        assert bad.status_code == 400

        app = await _submitted(applicant, event, answers={'annotation': 'Текст', 'work': file_id})
        assert app['files'][0]['name'] == 'работа.pdf'

        await login_as(db, stranger, email='stranger@example.org')
        assert (await stranger.get(f'/api/v1/files/{file_id}')).status_code == 404
        draft = (await stranger.post(f'/api/v1/events/{SLUG}/applications')).json()
        stolen = await stranger.put(
            f'/api/v1/applications/{draft["id"]}', json=_payload(event, answers={'annotation': 'x', 'work': file_id})
        )
        assert stolen.status_code == 422

    download = await client.get(f'/api/v1/files/{file_id}')
    assert download.status_code == 200 and download.content == b'%PDF-1.4 test'


async def test_direction_quota(db, client):
    event = await _open_event(db, client, limits={'per_direction': 1})
    async with _other(client) as first, _other(client) as second:
        await login_as(db, first, email='first@example.org')
        await _submitted(first, event)
        await login_as(db, second, email='second@example.org')
        app = (await second.post(f'/api/v1/events/{SLUG}/applications')).json()
        await second.put(f'/api/v1/applications/{app["id"]}', json=_payload(event))
        response = await second.post(f'/api/v1/applications/{app["id"]}/submit')
        assert response.json()['detail']['errors']['direction_id'] == 'В этом направлении больше нет мест'


async def test_form_validation(db, client):
    event = await _open_event(db, client)
    form = (await client.get(f'/api/v1/events/{event["id"]}/form')).json()
    form['fields'].append({'key': 'extra', 'type': 'text', 'label': 'Доп', 'visible_if': {'field': 'missing'}})
    assert (await client.put(f'/api/v1/events/{event["id"]}/form', json=form)).status_code == 422


async def test_registration_closed(db, client):
    event = await _open_event(db, client)
    await client.patch(
        f'/api/v1/events/{event["id"]}',
        json={'registration_closes_at': '2020-01-02T00:00:00Z', 'registration_opens_at': '2020-01-01T00:00:00Z'},
    )
    async with _other(client) as applicant:
        await login_as(db, applicant, email='late@example.org')
        response = await applicant.post(f'/api/v1/events/{SLUG}/applications')
        assert response.status_code == 422
        assert response.json()['detail']['message'] == 'Приём заявок сейчас закрыт'


async def test_export_and_bulk_decision(db, client):
    event = await _open_event(db, client)
    ids = []
    for email in ('a@example.org', 'b@example.org'):
        async with _other(client) as applicant:
            await login_as(db, applicant, email=email)
            ids.append((await _submitted(applicant, event))['id'])
    result = await client.post(
        f'/api/v1/events/{event["id"]}/applications/decision', json={'ids': ids, 'status': 'accepted'}
    )
    assert result.json() == {'updated': 2, 'skipped': []}

    export = await client.get(f'/api/v1/events/{event["id"]}/applications/export.xlsx')
    assert export.status_code == 200
    sheet = load_workbook(BytesIO(export.content)).active
    assert sheet['A1'].value == '№' and sheet.max_row == 3
    assert sheet['B2'].value == 'принята' and sheet['F2'].value == 'Иванова Анна'
