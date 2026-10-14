from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.modules.applications.form import ApplicationForm
from app.modules.applications.models import Application
from app.modules.applications.service import STATUS_TITLES


def _answer(value) -> str:
    if isinstance(value, bool):
        return 'да' if value else 'нет'
    if isinstance(value, list):
        return ', '.join(map(str, value))
    return '' if value is None else str(value)


def applications_xlsx(applications: list[Application], form: ApplicationForm) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Заявки'
    fields = [f for f in form.fields if f.type != 'file']
    header = [
        '№',
        'Статус',
        'Название',
        'Направление',
        'Форма участия',
        'Участники',
        'Классы',
        'Руководители',
        'Email заявителя',
        'Отправлена',
        *[f.label for f in fields],
        'Файлы',
        'Комментарий организаторов',
    ]
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='top')

    for app in applications:
        participants = [m for m in app.members if m.role == 'participant']
        supervisors = [m for m in app.members if m.role == 'supervisor']
        sheet.append(
            [
                app.number,
                STATUS_TITLES.get(app.status, app.status),
                app.title,
                app.direction.title if app.direction else '',
                app.participation_form.title if app.participation_form else '',
                '\n'.join(m.full_name for m in participants),
                '\n'.join(m.grade for m in participants if m.grade),
                '\n'.join(m.full_name for m in supervisors),
                app.user.email,
                app.submitted_at.strftime('%d.%m.%Y %H:%M') if app.submitted_at else '',
                *[_answer(app.answers.get(f.key)) for f in fields],
                '\n'.join(f.file.original_name for f in app.files),
                app.moderator_comment,
            ]
        )

    widths = {1: 6, 2: 16, 3: 40, 6: 32, 8: 28, 9: 28}
    for index in range(1, len(header) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(index, 20)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    sheet.freeze_panes = 'A2'

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
